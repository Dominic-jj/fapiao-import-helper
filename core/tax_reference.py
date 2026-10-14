"""税编参考表读取与规范化。"""

from __future__ import annotations

import os
import re

import openpyxl


def normalize_code(code) -> str:
    """把 Excel/JSON 中的税编统一成字符串。"""
    if code is None:
        return ""
    if isinstance(code, bool):
        return ""
    if isinstance(code, int):
        return str(code)
    if isinstance(code, float):
        if code.is_integer():
            return f"{code:.0f}"
        return str(code).strip()

    text = str(code).strip()
    if not text:
        return ""
    if re.fullmatch(r"\d+\.0+", text):
        return text.split(".", 1)[0]
    return text


def code_prefix(code: str) -> str:
    """税编表使用固定 19 位编码，尾部 0 前的部分代表层级前缀。"""
    return code.rstrip("0") or code


def mark_leaf_entries(entries: list[dict]) -> None:
    """按固定长度税编前缀标记叶子节点。"""
    prefix_rows = [
        (code_prefix(entry["code"]), entry["code"])
        for entry in entries
    ]

    for entry in entries:
        prefix = code_prefix(entry["code"])
        entry["is_leaf"] = not any(
            other_code != entry["code"]
            and len(other_prefix) > len(prefix)
            and other_prefix.startswith(prefix)
            for other_prefix, other_code in prefix_rows
        )


def load_tax_entries(path: str) -> list[dict]:
    """加载 tax_list.xlsx，返回统一字段的税编条目。"""
    if not os.path.exists(path):
        raise FileNotFoundError(f"税收分类表不存在: {path}")

    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = wb.active
    entries: list[dict] = []

    try:
        for row in ws.iter_rows(min_row=2, max_row=ws.max_row, values_only=True):
            # L列(index 11) = 合并编码, M列(12) = 名称, N列(13) = 分类名, O列(14) = 说明
            if len(row) < 13:
                continue

            code = row[11]
            name = row[12]

            if code is None or name is None:
                continue

            desc = row[14] if len(row) > 14 and row[14] else ""
            category = row[13] if len(row) > 13 and row[13] else ""
            normalized_code = normalize_code(code)
            if not normalized_code:
                continue

            entries.append(
                {
                    "code": normalized_code,
                    "name": str(name),
                    "category": str(category),
                    "desc": str(desc),
                }
            )
    finally:
        wb.close()

    entries.sort(key=lambda e: e["code"])
    mark_leaf_entries(entries)
    return entries
