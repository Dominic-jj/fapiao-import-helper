"""
Excel 生成模块 - 基于 invoice_template.xlsx 模板生成电子发票

读取模板文件，在第4行开始填入物品数据，生成新的 xlsx 文件。
"""

import io
import os

import openpyxl

from .config import TAX_RATE
from .models import compute_amount
from .paths import INVOICE_TEMPLATE_PATH

# 模板中承载明细的工作表名（不可随意更改）
DETAIL_SHEET_NAME = '1-明细模板'

_DEFAULT_TEMPLATE = INVOICE_TEMPLATE_PATH

# 模板中的列映射（Row 3 表头对应的列字母）
COLUMN_MAP = {
    'name': 'A',       # 项目名称
    'tax_code': 'B',   # 商品和服务税收分类编码
    'spec': 'C',       # 规格型号
    'unit': 'D',       # 单位
    'quantity': 'E',   # 商品数量
    'unit_price': 'F', # 商品单价
    'amount': 'G',     # 金额
    'tax_rate': 'H',   # 税率
    'discount': 'I',   # 折扣金额
    'policy': 'J',     # 优惠政策类型
    'coal': 'K',       # 煤炭种类
}

DATA_START_ROW = 4  # 数据从第4行开始（前3行是说明和表头）

# 对外暴露的税率常量统一来自 core/config.py
__all__ = [
    'COLUMN_MAP',
    'DATA_START_ROW',
    'DETAIL_SHEET_NAME',
    'TAX_RATE',
    'generate',
    'generate_to_bytes',
    'write_rows',
]


def _fill_rows(ws, items: list) -> None:
    """把明细写入工作表，并清除模板中可能残留的旧数据行。"""
    max_written_row = DATA_START_ROW + len(items) - 1
    if ws.max_row > max_written_row:
        ws.delete_rows(max_written_row + 1, ws.max_row - max_written_row)

    for i, item in enumerate(items):
        row = DATA_START_ROW + i
        # 列位置只仍由 COLUMN_MAP 定义，避免列字母在多处硬编码。
        values = {
            'name': item.name,
            'tax_code': str(item.tax_code) if item.tax_code else '',
            'spec': item.spec,
            'unit': item.unit,
            'quantity': item.quantity,
            'unit_price': item.unit_price,
            # 金额统一由系统重算并按 2 位小数写入，避免浮点尾数进入发票
            'amount': compute_amount(item.quantity, item.unit_price),
            'tax_rate': item.tax_rate or TAX_RATE,
            'discount': None,
            'policy': None,
            'coal': None,
        }
        for field, value in values.items():
            cell = ws[f'{COLUMN_MAP[field]}{row}']
            cell.value = value
            if isinstance(value, str):
                cell.data_type = 's'
        # 编码为纯数字字符串时，设为文本格式防止 Excel 显示科学计数法
        ws[f'{COLUMN_MAP["tax_code"]}{row}'].number_format = '@'


def write_rows(items: list, template_path: str = _DEFAULT_TEMPLATE) -> bytes:
    """基于模板生成 xlsx 字节流（不经过临时文件）。"""
    if not os.path.exists(template_path):
        raise FileNotFoundError(f'模板文件不存在: {template_path}')

    wb = openpyxl.load_workbook(template_path)
    try:
        ws = wb[DETAIL_SHEET_NAME]
        _fill_rows(ws, items)
        buffer = io.BytesIO()
        wb.save(buffer)
        return buffer.getvalue()
    finally:
        wb.close()


def generate(items: list, output_path: str,
             template_path: str = _DEFAULT_TEMPLATE) -> str:
    """
    生成电子发票 xlsx 文件。

    Args:
        items: parser.Item 列表（应已填充 tax_code）
        output_path: 输出文件路径
        template_path: 模板文件路径

    Returns:
        生成的文件路径
    """
    data = write_rows(items, template_path)
    parent = os.path.dirname(os.path.abspath(output_path))
    if parent:
        os.makedirs(parent, exist_ok=True)
    with open(output_path, 'wb') as f:
        f.write(data)

    return output_path


def generate_to_bytes(items: list,
                      template_path: str = _DEFAULT_TEMPLATE) -> bytes:
    """
    生成 xlsx 文件并返回字节数据。

    Args:
        items: parser.Item 列表
        template_path: 模板文件路径

    Returns:
        xlsx 文件的字节数据
    """
    return write_rows(items, template_path)
