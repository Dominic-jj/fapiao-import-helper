import shutil
import sys
import tempfile
import unittest
import json
from unittest.mock import patch
from pathlib import Path

import openpyxl


ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))

from core.tax_matcher import TaxMatcher  # noqa: E402


class TaxMatcherLeafCodeTest(unittest.TestCase):
    def setUp(self):
        self.temp_dir = Path(tempfile.mkdtemp())
        self.tax_list_path = self.temp_dir / "tax_list.xlsx"
        self.train_dir = self.temp_dir / "train_data"
        self.confirmed_map_path = self.temp_dir / "confirmed_tax_codes.json"
        self.train_dir.mkdir()
        self._write_tax_list()
        self.matcher = TaxMatcher(
            tax_list_path=str(self.tax_list_path),
            train_dir=str(self.train_dir),
            alias_map_path="",
            confirmed_map_path=str(self.confirmed_map_path),
        )

    def tearDown(self):
        shutil.rmtree(self.temp_dir)

    def test_corrupt_history_json_does_not_prevent_rule_matching(self):
        (self.train_dir / "_ground_truth.json").write_text("{broken", encoding="utf-8")
        matcher = TaxMatcher(str(self.tax_list_path), str(self.train_dir), "", str(self.confirmed_map_path))
        self.assertTrue(matcher.match_candidates("漏电开关"))

    def test_corrupt_confirmation_file_is_preserved_on_save(self):
        self.confirmed_map_path.write_text("{broken", encoding="utf-8")
        ok, message = self.matcher.save_confirmed_tax_code("漏电开关", "1090407040000000000")
        self.assertFalse(ok)
        self.assertIn("原文件已保留", message)
        self.assertEqual("{broken", self.confirmed_map_path.read_text(encoding="utf-8"))

    def test_invalid_confirmation_falls_back_to_valid_history(self):
        self.matcher.confirmed_map["漏电开关"] = "9999999999999999999"
        self.matcher.train_map["漏电开关"] = "1090407040000000000"
        self.assertEqual("confirmed", self.matcher.match_candidates("漏电开关")[0]["source"])

    def test_named_detail_sheet_is_used_when_active_sheet_is_other(self):
        wb = openpyxl.Workbook()
        ws = wb.create_sheet("1-明细模板")
        ws['A3'], ws['B3'] = '项目名称', '商品和服务税收分类编码'
        ws["A4"], ws["B4"] = "自定义商品", "1090407040000000000"
        wb.save(self.train_dir / "invoice.xlsx")
        wb.close()
        matcher = TaxMatcher(str(self.tax_list_path), str(self.train_dir), "", str(self.confirmed_map_path))
        self.assertEqual("1090407040000000000", matcher.train_map["自定义商品"])

    def test_frozen_history_preserves_conflicting_sources(self):
        bundled = self.temp_dir / "bundled"
        bundled.mkdir()
        (bundled / "_ground_truth.json").write_text(json.dumps({"old": [
            {"name": "内置商品", "tax_code": "1090407040000000000"},
            {"name": "覆盖商品", "tax_code": "1090407040000000000"},
        ]}), encoding="utf-8")
        (self.train_dir / "_ground_truth.json").write_text(json.dumps({"new": [
            {"name": "覆盖商品", "tax_code": "1090407050000000000"},
        ]}), encoding="utf-8")
        with patch.multiple("core.tax_matcher", BUNDLED_TRAIN_DATA_DIR=str(bundled),
                            TRAIN_DATA_DIR=str(self.train_dir), _DEFAULT_TRAIN_DIR=str(self.train_dir)):
            matcher = TaxMatcher(str(self.tax_list_path), str(self.train_dir), "", str(self.confirmed_map_path))
        self.assertEqual("1090407040000000000", matcher.match("内置商品")[0])
        self.assertEqual({'1090407040000000000', '1090407050000000000'},
                         {c['code'] for c in matcher.match_candidates('覆盖商品')})
        self.assertEqual(1, matcher.history_stats['conflicts'])

    def _write_tax_list(self):
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.append([""] * 15)

        def append_entry(code, name, category="", desc=""):
            row = [""] * 15
            row[11] = code
            row[12] = name
            row[13] = category
            row[14] = desc
            ws.append(row)

        append_entry("1000000000000000000", "货物", "货物")
        append_entry("1090000000000000000", "机械设备", "机械设备")
        append_entry("1090400000000000000", "开关设备", "电气设备")
        append_entry(
            "1090407040000000000",
            "配电或电器控制设备",
            "开关设备",
            "空气开关、漏电开关、控制开关",
        )
        append_entry(
            "1090407050000000000",
            "其他电气设备",
            "电气设备",
            "电机、控制柜",
        )
        append_entry(
            "1070601030000000000",
            "塑料管",
            "塑料制品",
            "塑料管、软管",
        )
        append_entry(
            "1080424000000000000",
            "金属制日用杂品",
            "金属制品",
            "水管、直通、弯头等金属管件",
        )
        wb.save(self.tax_list_path)
        wb.close()

    def test_fixed_width_parent_codes_are_not_leaf(self):
        self.assertFalse(self.matcher.is_leaf_code("1090400000000000000"))
        self.assertTrue(self.matcher.is_leaf_code("1090407040000000000"))

        is_valid, message = self.matcher.validate_tax_code("1090400000000000000")
        self.assertFalse(is_valid)
        self.assertIn("汇总商品", message)

        is_valid, message = self.matcher.validate_tax_code("1090407040000000000")
        self.assertTrue(is_valid)
        self.assertEqual("", message)

    def test_match_candidates_never_return_summary_code(self):
        candidates = self.matcher.match_candidates("空气开关", top_n=5)

        self.assertTrue(candidates)
        self.assertTrue(all(c["is_leaf"] for c in candidates))
        self.assertNotIn("1090400000000000000", {c["code"] for c in candidates})

    def test_confirmed_parent_is_resolved_to_leaf_descendant(self):
        self.matcher.train_map["空气开关"] = "1090400000000000000"

        candidates = self.matcher.match_candidates("空气开关", top_n=5)

        self.assertTrue(candidates)
        self.assertEqual("1090407040000000000", candidates[0]["code"])
        self.assertEqual("confirmed_parent", candidates[0]["source"])

    def test_alias_map_is_used_before_fuzzy_matching(self):
        alias_path = self.temp_dir / "tax_aliases.json"
        alias_path.write_text(
            json.dumps({"扎带": "1090400000000000000"}, ensure_ascii=False),
            encoding="utf-8",
        )
        matcher = TaxMatcher(
            tax_list_path=str(self.tax_list_path),
            train_dir=str(self.train_dir),
            alias_map_path=str(alias_path),
            confirmed_map_path=str(self.confirmed_map_path),
        )

        candidates = matcher.match_candidates("扎带", top_n=5)

        self.assertTrue(candidates)
        self.assertTrue(all(c["is_leaf"] for c in candidates))
        self.assertEqual("alias_parent", candidates[0]["source"])

    def test_alias_map_nested_groups_and_comment_keys(self):
        """分组嵌套 + `_` 前缀注释键应被展开，且注释不当作品项。"""
        alias_path = self.temp_dir / "tax_aliases.json"
        alias_path.write_text(
            json.dumps(
                {
                    "_说明": "这是一段分组说明，不应被当作映射。",
                    "_手动工具": {
                        "螺丝": "1090407040000000000",
                        "内六角扳手": {"tax_code": "1090407050000000000"},
                    },
                    "_空分组": {},
                },
                ensure_ascii=False,
            ),
            encoding="utf-8",
        )
        matcher = TaxMatcher(
            tax_list_path=str(self.tax_list_path),
            train_dir=str(self.train_dir),
            alias_map_path=str(alias_path),
            confirmed_map_path=str(self.confirmed_map_path),
        )

        # 注释键与空分组不污染映射
        self.assertNotIn("_说明", matcher.alias_map)
        self.assertNotIn("_手动工具", matcher.alias_map)
        self.assertNotIn("_空分组", matcher.alias_map)

        # 扁平 + 分组内的项都正确加载
        self.assertEqual("1090407040000000000", matcher.alias_map["螺丝"])
        self.assertEqual("1090407050000000000", matcher.alias_map["内六角扳手"])

        # 分组项也能匹配到叶子（此处映射的编码本身就是叶子 → source 为 alias）
        candidates = matcher.match_candidates("螺丝", top_n=5)
        self.assertTrue(candidates)
        self.assertTrue(all(c["is_leaf"] for c in candidates))
        self.assertEqual("alias", candidates[0]["source"])

    def test_confirmed_json_is_used_before_train_data(self):
        self.confirmed_map_path.write_text(
            json.dumps(
                {
                    "空气开关": {
                        "tax_code": "1090407040000000000",
                        "tax_name": "配电或电器控制设备",
                    }
                },
                ensure_ascii=False,
            ),
            encoding="utf-8",
        )
        matcher = TaxMatcher(
            tax_list_path=str(self.tax_list_path),
            train_dir=str(self.train_dir),
            alias_map_path="",
            confirmed_map_path=str(self.confirmed_map_path),
        )
        matcher.train_map["空气开关"] = "1090407050000000000"

        candidates = matcher.match_candidates("空气开关", top_n=5)

        self.assertTrue(candidates)
        self.assertEqual("1090407040000000000", candidates[0]["code"])
        self.assertEqual("confirmed", candidates[0]["source"])

    def test_save_confirmed_tax_code_writes_json_and_updates_memory(self):
        ok, message = self.matcher.save_confirmed_tax_code(
            "漏电开关",
            "1090407040000000000",
            "配电或电器控制设备",
        )

        self.assertTrue(ok, message)
        self.assertEqual("1090407040000000000", self.matcher.confirmed_map["漏电开关"])

        data = json.loads(self.confirmed_map_path.read_text(encoding="utf-8"))
        self.assertEqual("1090407040000000000", data["漏电开关"]["tax_code"])
        self.assertEqual("配电或电器控制设备", data["漏电开关"]["tax_name"])
        self.assertEqual("manual_confirmed", data["漏电开关"]["source"])

    def test_save_confirmed_tax_code_rejects_summary_code(self):
        ok, message = self.matcher.save_confirmed_tax_code(
            "空气开关",
            "1090400000000000000",
            "开关设备",
        )

        self.assertFalse(ok)
        self.assertIn("汇总商品", message)


if __name__ == "__main__":
    unittest.main()
